import io
from types import SimpleNamespace

import openpyxl
import pypdf
from docx import Document
from PIL import Image

from memoria import file_extraction


def _pdf_senza_testo() -> bytes:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _pdf_immagine() -> bytes:
    """PDF "scansione": una pagina che è solo un'immagine, zero testo -
    esattamente quello che produce uno scanner."""
    buffer = io.BytesIO()
    Image.new("RGB", (80, 80), color=(200, 200, 200)).save(buffer, format="PDF")
    return buffer.getvalue()


def _pdf_testo(testo: str = "Fattura n. 123 del 16/07/2026 - Importo dovuto: 500 EUR") -> bytes:
    """PDF minimale con vero strato di testo, costruito a mano (pypdf non
    sa scrivere testo e non vogliamo reportlab solo per i test)."""
    contenuto = f"BT /F1 12 Tf 72 720 Td ({testo}) Tj ET".encode("latin-1")
    oggetti = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(contenuto)).encode() + b" >>\nstream\n" + contenuto + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, corpo in enumerate(oggetti, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode())
        out.write(corpo)
        out.write(b"\nendobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(oggetti) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(oggetti) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
    )
    return out.getvalue()


def _pdf_misto() -> bytes:
    """Trappola reale: contratto con copertina digitale + pagine scansionate.
    Il routing per soglia sul testo TOTALE lo classificava "digitale" e
    perdeva in silenzio le pagine scansionate."""
    writer = pypdf.PdfWriter()
    writer.append(pypdf.PdfReader(io.BytesIO(_pdf_testo())))
    writer.append(pypdf.PdfReader(io.BytesIO(_pdf_immagine())))
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _pdf_cifrato() -> bytes:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("segreta")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_pdf_cifrato_rilevato():
    assert file_extraction.pdf_e_cifrato(_pdf_cifrato()) is True
    assert file_extraction.pdf_e_cifrato(_pdf_testo()) is False


def test_pdf_corrotto_non_e_cifrato():
    """Un file illeggibile non è "cifrato": deve seguire il percorso normale
    (visione), che dà un errore esplicito, non il messaggio sulla password."""
    assert file_extraction.pdf_e_cifrato(b"non sono un PDF") is False


def test_indici_pagine_scansione_pdf_misto():
    assert file_extraction.indici_pagine_scansione(_pdf_misto()) == [1]


def test_indici_pagine_scansione_pdf_solo_testo():
    assert file_extraction.indici_pagine_scansione(_pdf_testo()) == []


def test_pagina_bianca_senza_immagini_non_e_scansione():
    """Una pagina bianca (niente testo, niente immagini) è comune nei
    documenti digitali - non deve far scattare il costoso percorso visivo."""
    assert file_extraction.indici_pagine_scansione(_pdf_senza_testo()) == []


def test_indici_pagine_scansione_pdf_corrotto_non_crash():
    assert file_extraction.indici_pagine_scansione(b"non sono un PDF") == []


def test_pdf_senza_testo_non_ha_testo_digitale():
    assert file_extraction.pdf_ha_testo_digitale(_pdf_senza_testo()) is False


def test_pdf_con_testo_digitale_rilevato(monkeypatch):
    """Trappola: pdf_ha_testo_digitale deve riconoscere un PDF con testo
    reale (non solo il caso vuoto) - pypdf.PdfReader mockato per non
    dipendere da una libreria di generazione PDF con testo (reportlab)."""
    pagina_finta = SimpleNamespace(extract_text=lambda: "Fattura n. 123 del 16/07/2026 - Importo dovuto: 500 EUR")
    lettore_finto = SimpleNamespace(pages=[pagina_finta])
    monkeypatch.setattr(file_extraction.pypdf, "PdfReader", lambda _: lettore_finto)

    assert file_extraction.pdf_ha_testo_digitale(b"non importa, mockato") is True
    assert "Fattura n. 123" in file_extraction.estrai_testo_pdf(b"non importa, mockato")


def test_pdf_corrotto_trattato_come_senza_testo_non_crash():
    assert file_extraction.pdf_ha_testo_digitale(b"non sono affatto un PDF") is False


def test_numero_pagine_pdf():
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    assert file_extraction.numero_pagine_pdf(buffer.getvalue()) == 2


def test_estrai_testo_docx():
    documento = Document()
    documento.add_paragraph("Contratto di fornitura con Rossi Srl")
    buffer = io.BytesIO()
    documento.save(buffer)

    testo = file_extraction.estrai_testo_docx(buffer.getvalue())

    assert "Contratto di fornitura con Rossi Srl" in testo


def test_estrai_testo_xlsx():
    cartella = openpyxl.Workbook()
    foglio = cartella.active
    foglio.title = "Listino"
    foglio.append(["Prodotto", "Prezzo"])
    foglio.append(["Vite M6", 0.15])
    buffer = io.BytesIO()
    cartella.save(buffer)

    testo = file_extraction.estrai_testo_xlsx(buffer.getvalue())

    assert "Listino" in testo
    assert "Vite M6" in testo
